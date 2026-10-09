"""Presentation built from the author's actual slide templates.

The reference is copied verbatim into design/presentation_reference.pptx.
Slide furniture and reused layouts retain their original OOXML typography.
New plots consist of editable PowerPoint shapes, except the spatial hexbin.
"""
from copy import deepcopy
import json
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
ORANGE='E8591A';DARK='1E2227';PAPER='F2EFE8';GRAY='5B636E';BORDER='C9C3B6';MUTED='B9B4A8'


def n(v):return f'{int(v):,}'.replace(',','\u00a0')
def ru(v,d=1):return f'{v:.{d}f}'.replace('.',',')


def set_text(shape,value):
    """Replace text while retaining paragraph/rPr formatting from the reference."""
    frame=shape.text_frame
    prototypes=[deepcopy(p._p) for p in frame.paragraphs]
    frame.clear()
    for i,line in enumerate(str(value).split('\n')):
        proto=deepcopy(prototypes[min(i,len(prototypes)-1)])
        runs=proto.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}r')
        props=deepcopy(runs[0].find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')) if runs else None
        for node in list(proto):
            if node.tag.rsplit('}',1)[-1] in ['r','br','fld']:proto.remove(node)
        r=OxmlElement('a:r')
        if props is not None:r.append(props)
        txt=OxmlElement('a:t');txt.text=line;r.append(txt);proto.append(r)
        if i==0:
            old=frame.paragraphs[0]._p;old.getparent().replace(old,proto)
        else:frame._txBody.append(proto)


def clone_shape(slide,shape):
    el=deepcopy(shape.element);slide.shapes._spTree.insert_element_before(el,'p:extLst')
    return slide.shapes[-1]


def t(slide,x,y,w,h,value,size=12,font='Arial',color=DARK,bold=False,align=PP_ALIGN.LEFT,spacing=0,line=1.0,center=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE if center else MSO_ANCHOR.TOP
    for i,line_text in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.line_spacing=line
        p.space_before=p.space_after=Pt(0)
        r=p.add_run();r.text=line_text;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if spacing:r._r.get_or_add_rPr().set('spc',str(spacing))
    return sh


def box(s,x,y,w,h,fill='FFFFFF',stroke=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    style=sh.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:sh.element.remove(style)
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.6)
    else:sh.line.fill.background()
    return sh


def rule(s,x,y,w,color=BORDER):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x+w),Inches(y))
    style=sh.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:sh.element.remove(style)
    sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(.5)
    return sh


def label(s,x,y,w,value,color=GRAY):
    return t(s,x,y,w,.26,value,10,'Consolas',color,True,spacing=150)


def table(s,headers,rows,widths,y=2.0,row_h=.6,fonts=None,highlight=()):
    """Dark header, alternating paper rows and mono values, as reference slide 9."""
    box(s,.7,y,11.9,.43,DARK)
    x=.7
    for h,w in zip(headers,widths):t(s,x+.15,y,w-.3,.43,h,9,'Consolas','FFFFFF',True,center=True,spacing=100);x+=w
    y+=.43
    for i,row in enumerate(rows):
        if i%2:box(s,.7,y,11.9,row_h,'E8E4DB')
        x=.7
        for j,(val,w) in enumerate(zip(row,widths)):
            font,size,bold=('Arial',12,False) if fonts is None else fonts[j]
            color=ORANGE if (i,j) in highlight else DARK
            t(s,x+.15,y+.04,w-.3,row_h-.08,val,size,font,color,bold,center=True,line=1.05)
            x+=w
        rule(s,.7,y+row_h,11.9);y+=row_h
    return y


def card(s,x,y,w,title,status,body,accent=True,h=.95):
    box(s,x,y,w,h,'FFFFFF',BORDER);box(s,x,y,.08,h,ORANGE if accent else BORDER)
    t(s,x+.25,y+.12,w*.48,.4,title,17,'Bahnschrift',DARK,True)
    t(s,x+w*.50,y+.16,w*.46,.3,status,9,'Consolas',ORANGE if accent else GRAY,True,align=PP_ALIGN.RIGHT,spacing=60)
    t(s,x+.25,y+.55,w-.45,.35,body,10.5,'Arial',GRAY)


def panel(s,value,subtitle,facts,y=2.0,h=3.6):
    box(s,9.95,y,2.65,h,DARK)
    t(s,10.2,y+.2,2.15,.85,value,30 if len(value)>6 else 42,'Bahnschrift',ORANGE,True)
    t(s,10.2,y+1.05,2.15,.65,subtitle,12,'Arial','FFFFFF',True,line=1.05)
    yy=y+1.87
    for fact in facts:
        box(s,10.2,yy+.06,.08,.08,ORANGE)
        t(s,10.4,yy,1.95,.56,fact,10,'Arial',MUTED,line=1.0)
        yy+=.62


def horizontal(s,labels,values,fmt=lambda x:ru(x),maximum=None,y=2.35,h=3.3,color=ORANGE):
    maximum=maximum if maximum is not None else max(values)
    step=h/len(labels)
    for i,(name,value) in enumerate(zip(labels,values)):
        yy=y+i*step
        t(s,.7,yy,3.0,step*.70,name,10,'Consolas',GRAY,center=True)
        box(s,3.85,yy+.07,4.7,step*.47,PAPER,BORDER)
        if value>0:box(s,3.85,yy+.07,max(.009,4.7*value/maximum),step*.47,color)
        t(s,8.7,yy,1.0,step*.72,fmt(value),12,'Consolas',color,True,align=PP_ALIGN.RIGHT,center=True)


def vertical(s,labels,values,y=2.4,h=2.65,x=.9,w=8.5,colors=None,show_values=False):
    maximum=max(values)*1.12
    for v in [0,.25,.5,.75,1]:
        yy=y+h*(1-v);rule(s,x,yy,w)
        t(s,x-.2,yy-.16,.8,.22,n(maximum*v),8,'Consolas',GRAY,align=PP_ALIGN.RIGHT)
    left=x+.75;plotw=w-.75;step=plotw/len(labels)
    for i,(name,value) in enumerate(zip(labels,values)):
        color=colors[i] if colors else DARK;barh=h*value/maximum;xx=left+i*step+.045
        box(s,xx,y+h-barh,max(.04,step-.09),barh,color)
        t(s,left+i*step,y+h+.12,step,.28,str(name),8 if len(labels)>15 else 10,'Consolas',GRAY,align=PP_ALIGN.CENTER)
        if show_values:t(s,left+i*step,y+h-barh-.30,step,.25,n(value),9,'Consolas',color,True,align=PP_ALIGN.CENTER)


class Template:
    def __init__(self):
        self.ref=Presentation(ROOT/'design/presentation_reference.pptx')
        self.deck=Presentation();self.deck.slide_width=self.ref.slide_width;self.deck.slide_height=self.ref.slide_height

    def page(self,section,title,source=None,historical=False):
        slide=self.deck.slides.add_slide(self.deck.slide_layouts[6])
        src=self.ref.slides[1 if source is None else source]
        for sh in list(src.shapes)[:16] if source is None else src.shapes:clone_shape(slide,sh)
        set_text(slide.shapes[3],section);set_text(slide.shapes[4],title)
        set_text(slide.shapes[6], 'Chicago Building Violations · замеры 26.09.2026 12:17' if historical else 'Chicago Building Violations · анализ 09.10.2026')
        set_text(slide.shapes[13],str(len(self.deck.slides)))
        return slide

    def cover(self,q):
        s=self.deck.slides.add_slide(self.deck.slide_layouts[6])
        for sh in self.ref.slides[0].shapes:clone_shape(s,sh)
        edits={'ЗАДАЧА 4 · РК-1 · DATA SCIENCE':'КУРСОВОЙ ПРОЕКТ · РК-1 · DATA SCIENCE',
               'Обоснование программных средств на полном архиве\npandas  ·  Parquet  ·  DuckDB':'Аналитическая часть: качество данных, EDA и baseline\npandas  ·  Parquet  ·  DuckDB  ·  scikit-learn',
               '2\u00a0030\u00a0101 запись\n930 МБ в CSV':f"{n(q['rows'])} записей\n{q['csv_mb']:.0f} МБ в CSV",
               'ДАТА ЗАМЕРОВ':'ДАТА АНАЛИЗА','26.09.2026 12:17':'09.10.2026'}
        for sh in s.shapes:
            if sh.has_text_frame and sh.text in edits:set_text(sh,edits[sh.text])
        return s

    def save(self):
        total=len(self.deck.slides)
        for s in list(self.deck.slides)[1:]:set_text(s.shapes[15],str(total))
        self.deck.save(ROOT/'presentation.pptx')
        self.deck.save(ROOT/'presentation_medikhanov_timur.pptx')
        return total


def make_presentation(q,m,eda,checks):
    tpl=Template();tpl.cover(q)
    d=pd.read_parquet(ROOT/'data/processed/inspections.parquet')
    df=pd.read_parquet(ROOT/'data/processed/model_data.parquet')
    test=[r for r in m['metrics'] if r['split']=='test'];lr=next(r for r in test if r['model']=='Logistic regression');dummy=next(r for r in test if r['model']=='Dummy prior')
    # 2: user, action, error prices and success criterion, with the source's card layout.
    s=tpl.page('01 — ПОСТАНОВКА ЗАДАЧИ','Риск инспекции: пользователь и решение')
    label(s,.7,2.0,6.1,'КОНТЕКСТ СТРОЙКОНТРОЛЯ')
    t(s,.7,2.35,6.1,1.40,'Координатор инспекций распределяет внимание между уже запланированными проверками. Модель оценивает вероятность FAILED до осмотра; инспектор подтверждает фактический результат.\nЕдиница анализа — одна инспекция объекта.',15,line=1.15)
    card(s,7.2,2.0,5.4,'Пользователь','КООРДИНАТОР','Руководитель службы строительного контроля')
    card(s,7.2,3.08,5.4,'Решение','ПРИОРИТЕТ','Какие проверки требуют более раннего внимания')
    card(s,7.2,4.16,5.4,'Результат','ВЕРОЯТНОСТЬ','P(FAILED | служба, категория, дата, координаты)')
    card(s,7.2,5.24,5.4,'Критерий','TEST 2025','AP ≥ Dummy + 0,05; recall ≥ 0,80; precision ≥ доли FAILED')
    label(s,.7,4.10,6.1,'ЦЕНА ОШИБКИ · УЧЕБНАЯ ГИПОТЕЗА 5:1')
    t(s,.7,4.49,1.0,.65,'FN',30,'Bahnschrift',ORANGE,True)
    t(s,1.9,4.48,4.7,.65,'5 у.е. — пропущенная проблемная проверка; задержка внимания к нарушениям.',13)
    t(s,.7,5.30,1.0,.65,'FP',30,'Bahnschrift',DARK,True)
    t(s,1.9,5.30,4.7,.65,'1 у.е. — лишнее внимание к успешной проверке; расход времени специалиста.',13)
    t(s,.7,6.20,11.9,.28,'Стоимость не измерена в деньгах; доступность входных полей до осмотра требует подтверждения.',11,'Consolas',GRAY)
    # 3: retain the exact source passport, process strip and status bar geometry.
    s=tpl.page('02 — ЗАДАЧА И ДАННЫЕ','Контекст и структура данных',source=1)
    set_text(s.shapes[17],'Каждая инспекция объекта может содержать несколько записей нарушения. Аналитика объединяет этот поток на уровне проверки и оценивает риск её результата. ERP-приложение — дальнейшая часть проекта; здесь оценивается аналитика.')
    for idx,value in {34:n(q['rows']),40:'01.01.2006 – 08.10.2026',43:n(q['objects']),46:n(q['violation_types']),49:n(q['inspection_ids'])}.items():set_text(s.shapes[idx],value)
    raw=q['raw_status_counts'];total=q['rows'];xx=.7
    for idx,status in zip([55,56,57],['FAILED','PASSED','CLOSED']):
        sh=s.shapes[idx];sh.left=Inches(xx);sh.width=Inches(11.933*raw[status]/total);xx+=sh.width/914400
    for idx,status in zip([59,61,63],['FAILED','PASSED','CLOSED']):set_text(s.shapes[idx],f"{status} {ru(raw[status]/total*100)} %")
    # 4: restore sample-to-full-archive comparison from the reference.
    s=tpl.page('03 — МАСШТАБ','От выборки к полному архиву',source=2)
    edits={'2\u00a0030\u00a0101':n(q['rows']),'930,39 МБ':ru(q['csv_mb'],2)+' МБ','01.01.2006 – 25.09.2026':'01.01.2006 – 08.10.2026'}
    for sh in s.shapes:
        if sh.has_text_frame and sh.text in edits:set_text(sh,edits[sh.text])
    # 5: exact original data-layer cards, orange spines, arrows and typefaces.
    s=tpl.page('04 — СТРУКТУРА ПРОЕКТА','Слои данных: raw → processes → features',source=3)
    for idx,value in {20:'data/raw/\nCSV + Parquet ZSTD',22:n(q['rows']),29:'data/processed/\ninspections.parquet',31:n(q['aggregated_inspections']),38:'data/processed/\nmodel_data.parquet',40:'8',41:'входов · цель target_failed',42:'src/pipeline.py — полный расчёт · notebooks/course_project.ipynb — анализ · src/presentation.py — сборка слайдов'}.items():set_text(s.shapes[idx],value)
    # 6: restore a detailed seven-row preparation ledger.
    s=tpl.page('05 — ПРЕДОБРАБОТКА','Что сделано с данными в processes/')
    table(s,['№','ШАГ','ЧТО СДЕЛАНО','ЭФФЕКТ'],[
        ['01','Дубликаты','Повторы ID → последняя версия по дате изменения',n(q['duplicate_ids'])],
        ['02','Категории','strip + upper; пустые строки → NULL','4 категориальных поля'],
        ['03','Даты','Явный разбор; диапазон 2006 — 09.10.2026','0 ошибок разбора'],
        ['04','Координаты','Пустые / вне рамки → NULL; инспекция сохраняется',n(q['missing_coordinates'])+' строк'],
        ['05','Агрегация','Группировка по inspection_number; проверка атрибутов',n(q['aggregated_inspections'])+' инспекций'],
        ['06','Дата осмотра','Разные violation_date внутри номера → неоднозначность',n(q['ambiguous_inspections'])+' исключений'],
        ['07','Метка','FAILED / PASSED; остальные статусы исключены',n(q['labeled_inspections'])+' для модели']],
        [.65,2.2,5.7,3.35],row_h=.50,fonts=[('Consolas',12,True),('Bahnschrift',14,True),('Arial',12,False),('Consolas',12,True)],highlight=[(i,0) for i in range(7)]+[(3,3),(5,3)])
    t(s,.7,6.16,11.9,.30,'41 744 — неоднозначность даты, а не доказанная ошибка источника. Без противоречий статуса, объекта и службы.',10.5,'Consolas',GRAY)
    # 7: no hidden simplification of the former fifteen-feature setup.
    s=tpl.page('06 — ПРИЗНАКИ','Какие поля доступны до текущей инспекции')
    label(s,.7,2.0,6.1,'ВХОДЫ ПРОСТОЙ МОДЕЛИ · 8 ПОЛЕЙ')
    t(s,.7,2.4,6.1,3.95,'department_bureau — назначенная служба\ninspection_category — категория проверки\nyear — календарный год\nmonth_sin / month_cos — циклический месяц\ndayofweek — день недели\nlatitude / longitude — координаты объекта\n\nВ прежнем notebook было 15 признаков, включая историю результатов. История исключена из baseline: время доступности результатов и порядок проверок одного дня не подтверждены.',14,'Arial',DARK,line=1.18)
    card(s,7.2,2.0,5.4,'Категории','ONE-HOT','Неизвестная категория на тесте → handle_unknown=ignore')
    card(s,7.2,3.08,5.4,'Числа','TRAIN ONLY','Медиана + индикатор пропуска + StandardScaler')
    card(s,7.2,4.16,5.4,'Постфактум','ИСКЛЮЧЕНО','n_violations, n_codes, текущий статус, устранение, waived')
    card(s,7.2,5.24,5.4,'История','СЛЕДУЮЩИЙ ЭТАП','Вернуть после проверки доступности событий во времени')
    # 8–9: complete, quantified eight-item audit, dark table header.
    for first in [0,4]:
        s=tpl.page(f'07 — КАЧЕСТВО ДАННЫХ · {first+1}–{first+4}','Проверки качества: результат и решение')
        table(s,['ПРОВЕРКА','ИЗМЕРЕННЫЙ РЕЗУЛЬТАТ','РЕШЕНИЕ'],checks[first:first+4],[2.6,4.1,5.2],row_h=.78,fonts=[('Bahnschrift',14,True),('Consolas',11,True),('Arial',12,False)])
        t(s,.7,6.0,11.9,.45,'Стандартный чек-лист: подробный перечень преподавателя не предоставлен.\nПолный raw-аудит без подвыборки; ошибки чтения строк не игнорируются.',11,'Consolas',GRAY,line=1.1)
    # 10–16: seven editable analytical plots in the source's bar + black-panel layout.
    missing=pd.read_csv(ROOT/'reports/missingness.csv').sort_values('missing_pct',ascending=False).head(8)
    s=tpl.page('08 — EDA · 1/7 · ПОЛНОТА','Пропуски в исходном архиве нарушений')
    label(s,.7,2.0,8.8,'ДОЛЯ NULL И ПУСТЫХ СТРОК · RAW · %')
    horizontal(s,missing.column.tolist(),missing.missing_pct.tolist(),fmt=lambda v:ru(v,2)+' %',maximum=100,y=2.45,h=3.0)
    panel(s,'80,6 %','пропусков в поле SSA',['Не используется в baseline','Координаты: медиана train + индикатор','Даты статуса и комментарии исключены'],h=4.20)
    t(s,.7,5.75,8.8,.7,eda[0]['conclusion'],13,line=1.1)
    annual=d.groupby(d.inspection_date.dt.year).size()
    s=tpl.page('08 — EDA · 2/7 · ВРЕМЯ','Количество инспекций по годам')
    label(s,.7,2.0,8.8,'ОДНА ИНСПЕКЦИЯ = ОДИН НОМЕР · 2006–2026')
    vertical(s,annual.index,annual.values,colors=[ORANGE if y==2026 else DARK for y in annual.index],y=2.55)
    panel(s,n(annual.max()),'максимум: 2006 год',['2026 — неполный год','Объём зависит от регистрации','Test модели: полный 2025 год'],h=4.20)
    t(s,.7,5.75,8.8,.7,eda[1]['conclusion'],13,line=1.1)
    counts=d.inspection_status.fillna('UNKNOWN').value_counts()
    s=tpl.page('08 — EDA · 3/7 · ЦЕЛЬ','Статусы на уровне инспекции')
    label(s,.7,2.0,8.8,'ДО ФИЛЬТРАЦИИ СТАТУСОВ И НЕОДНОЗНАЧНЫХ ДАТ')
    vertical(s,counts.index,counts.values,colors=[ORANGE if v=='FAILED' else DARK for v in counts.index],show_values=True,y=2.55)
    panel(s,'73,7 %','FAILED в бинарной выборке',['FAILED=1; PASSED=0','CLOSED / HOLD / UNKNOWN исключены','Accuracy дополняется метриками классов'],h=4.20)
    t(s,.7,5.75,8.8,.70,eda[2]['conclusion'],13,line=1.1)
    bins=pd.cut(d.n_violations,[0,1,2,5,10,20,np.inf],labels=['1','2','3–5','6–10','11–20','>20']);hist=bins.value_counts(sort=False)
    s=tpl.page('08 — EDA · 4/7 · ГРАНУЛЯРНОСТЬ','Сколько нарушений приходится на инспекцию')
    label(s,.7,2.0,8.8,'ГРУППЫ ПО ЧИСЛУ ЗАПИСЕЙ НАРУШЕНИЯ')
    vertical(s,hist.index.astype(str),hist.values,y=2.55,show_values=True)
    panel(s,'3','записи: медиана',['95-й перцентиль — 13','Максимум — 538 записей','Текущие нарушения используются только в EDA'],h=4.20)
    t(s,.7,5.75,8.8,.70,eda[3]['conclusion'],13,line=1.1)
    bureau=df.groupby('department_bureau').target_failed.agg(['size','mean']).sort_values('size',ascending=False).head(10).sort_values('mean',ascending=False)
    s=tpl.page('08 — EDA · 5/7 · СЛУЖБЫ','Доля FAILED в десяти крупнейших службах')
    label(s,.7,2.0,8.8,'БИНАРНАЯ ВЫБОРКА · 95 % ИНТЕРВАЛЫ УИЛСОНА')
    horizontal(s,bureau.index.tolist(),(bureau['mean']*100).tolist(),fmt=lambda v:ru(v)+' %',maximum=100,y=2.42,h=3.15)
    for i,(_,r) in enumerate(bureau.iterrows()):
        p=r['mean'];nn=r['size'];z=1.96;center=(p+z*z/(2*nn))/(1+z*z/nn);half=z*np.sqrt(p*(1-p)/nn+z*z/(4*nn*nn))/(1+z*z/nn)
        yy=2.42+i*.315+.145;rule(s,3.85+4.7*(center-half),yy,4.7*(2*half),DARK)
        t(s,7.70,yy-.25,.8,.15,'n='+n(nn),7,'Consolas',GRAY)
    panel(s,'31–93 %','разброс долей FAILED',['Процедуры учёта различаются','Интервалы описательные','Повторы объектов не учитываются'],h=4.20)
    t(s,.7,5.80,8.8,.65,eda[4]['conclusion'],12.5,line=1.1)
    season=df[df.year<=2025].groupby('month').target_failed.agg(['size','mean'])
    s=tpl.page('08 — EDA · 6/7 · СЕЗОННОСТЬ','Месяц: объём проверок и доля FAILED')
    label(s,.7,2.0,4.0,'ЧИСЛО ИНСПЕКЦИЙ · 2006–2025')
    label(s,5.45,2.0,4.0,'ДОЛЯ FAILED · %')
    vertical(s,season.index,season['size'].tolist(),x=.8,w=4.05,y=2.65,h=2.6)
    x0=5.65;yy=2.65;hh=2.6;ww=3.7
    for pc in [0,25,50,75,100]:
        rule(s,x0,yy+hh*(1-pc/100),ww);t(s,x0-.45,yy+hh*(1-pc/100)-.12,.35,.24,str(pc),8,'Consolas',GRAY,align=PP_ALIGN.RIGHT)
    prev=None
    for i,(mo,r) in enumerate(season.iterrows()):
        xx=x0+i*ww/11;py=yy+hh*(1-r['mean']);box(s,xx-.035,py-.035,.07,.07,ORANGE)
        if prev:
            sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(prev[0]),Inches(prev[1]),Inches(xx),Inches(py));sh.line.color.rgb=RGBColor.from_string(ORANGE);sh.line.width=Pt(1.5)
        prev=(xx,py);t(s,xx-.12,yy+hh+.15,.24,.22,str(mo),8,'Consolas',GRAY,align=PP_ALIGN.CENTER)
    panel(s,'72–75 %','месячная доля FAILED',['Только полные годы','Месяц кодируется sin / cos','Сезон смешан с составом служб'],h=4.20)
    t(s,.7,5.82,8.8,.65,eda[5]['conclusion'],13,line=1.1)
    s=tpl.page('08 — EDA · 7/7 · ГЕОГРАФИЯ','Где концентрируются зарегистрированные проверки')
    label(s,.7,2.0,8.8,'ПЛОТНОСТЬ ИНСПЕКЦИЙ · ЛОГАРИФМИЧЕСКАЯ ШКАЛА')
    # Dense spatial plot is the sole raster image; geometry and insight remain native.
    from PIL import Image
    image=Image.open(ROOT/'reports/figures/eda07_geo.png')
    # Remove its duplicate matplotlib title; native slide supplies the title and caption.
    crop=image.crop((0,int(image.height*.13),image.width,image.height))
    path=ROOT/'reports/figures/presentation_geo.png';crop.save(path)
    s.shapes.add_picture(str(path),Inches(.7),Inches(2.22),width=Inches(8.95))
    panel(s,'502','инспекции без координат',['Карта показывает наблюдение','Нет нормировки на число зданий','Нельзя оценить опасность района'],h=4.20)
    t(s,.733,6.03,8.8,.55,eda[6]['conclusion'],12.5,line=1.1)
    # 17: precise learning/threshold/test protocol.
    s=tpl.page('09 — BASELINE · ПРОТОКОЛ','Временное разделение и обучение модели')
    table(s,['ВЫБОРКА','ПЕРИОД','ИНСПЕКЦИИ','FAILED'],[[name,v['date_from']+' — '+v['date_to'],n(v['n']),ru(v['failed_share']*100)+' %'] for name,v in m['splits'].items()],[2.0,4.3,3.0,2.6],row_h=.58,fonts=[('Consolas',12,True),('Consolas',12,True),('Consolas',14,True),('Consolas',14,True)])
    box(s,.7,4.38,5.70,1.70,'FFFFFF',BORDER);box(s,.7,4.38,.08,1.7,ORANGE)
    label(s,1.0,4.60,5.2,'DUMMYCLASSIFIER · STRATEGY=PRIOR',ORANGE)
    t(s,1.0,5.0,5.05,.85,'Одна вероятность для всех инспекций — частота FAILED по train. Показывает, что можно получить без использования признаков.',13,line=1.10)
    box(s,6.9,4.38,5.70,1.70,'FFFFFF',BORDER);box(s,6.9,4.38,.08,1.7,ORANGE)
    label(s,7.2,4.60,5.2,'LOGISTIC REGRESSION · L2 · C=1',ORANGE)
    t(s,7.2,5.0,5.05,.85,'Pipeline: OneHotEncoder + медиана + StandardScaler. lbfgs; max_iter=2000. Все преобразования обучаются только на train.',13,line=1.10)
    t(s,.7,6.25,11.9,.27,'Порог: минимум 5×FN+FP на 2024. Test 2025 не участвует в настройке; 2026 неполный и исключён.',10.5,'Consolas',GRAY)
    # 18: keep all metrics, sample sizes and costs rather than a headline AP alone.
    s=tpl.page('09 — BASELINE · МЕТРИКИ','Две простые модели: проверка на 2025 годе')
    keys=[('Average Precision','average_precision'),('ROC-AUC','roc_auc'),('Precision FAILED','precision'),('Recall FAILED','recall'),('Balanced accuracy','balanced_accuracy'),('F1 FAILED','f1'),('Brier (меньше лучше)','brier'),('Стоимость (5×FN+FP)/N','cost_per_inspection')]
    table(s,['МЕТРИКА','DUMMY PRIOR','LOGISTIC REGRESSION'],[[name,ru(dummy[k],3),ru(lr[k],3)] for name,k in keys],[5.3,3.2,3.4],row_h=.41,fonts=[('Arial',13,False),('Consolas',14,True),('Consolas',14,True)],highlight=[(i,2) for i in [0,1,7]])
    t(s,.7,5.95,11.9,.55,f"Test: {n(lr['n'])} инспекций; FAILED {ru(lr['prevalence']*100)} %. Порог: {ru(lr['threshold'],3)}.\nAP — average_precision_score, без трапецеидальной интерполяции PR-кривой.",11,'Consolas',GRAY,line=1.1)
    # 19: native matrix cells, a useful explicit explanation of the cost tradeoff.
    s=tpl.page('09 — BASELINE · ОШИБКИ','Что стоит за высоким recall')
    for x,name,r in [(.7,'Dummy prior',dummy),(5.3,'Logistic regression',lr)]:
        label(s,x,2.0,4.15,name.upper())
        t(s,x+.80,2.50,1.3,.30,'PASSED',10,'Consolas',GRAY,True,align=PP_ALIGN.CENTER)
        t(s,x+2.20,2.50,1.3,.30,'FAILED',10,'Consolas',GRAY,True,align=PP_ALIGN.CENTER)
        for i,(actual,vals) in enumerate([('PASSED',[r['tn'],r['fp']]),('FAILED',[r['fn'],r['tp']])]):
            t(s,x,2.95+i*1.02,.75,.85,actual,9,'Consolas',GRAY,center=True)
            for j,value in enumerate(vals):
                xx=x+.8+j*1.40;yy=2.95+i*1.02
                box(s,xx,yy,1.30,.90,ORANGE if (i==1 and j==1) else 'FFFFFF',BORDER)
                t(s,xx,yy,1.30,.90,n(value),24,'Bahnschrift','FFFFFF' if (i==1 and j==1) else DARK,True,align=PP_ALIGN.CENTER,center=True)
        t(s,x+.8,5.20,2.70,.30,'строки — факт; столбцы — прогноз',9,'Consolas',GRAY,align=PP_ALIGN.CENTER)
    panel(s,'97,2 %','проверок получают приоритет',['FN: 20; FP: 1 968','Ложные тревоги у PASSED: 87,5 %','Условная стоимость ниже Dummy на 8,0 %'],h=4.20)
    t(s,.7,5.88,8.8,.60,'Высокий recall достигнут широкой очередью. Учебный критерий выполнен, но экономия внимания ещё не доказана: нужен лимит очереди и precision@K.',13,line=1.10)
    # 20–21: full plan and risk mitigation, not four abbreviated labels.
    s=tpl.page('10 — ДАЛЬНЕЙШАЯ РАБОТА','План после первого рубежного контроля')
    table(s,['ЭТАП','РАБОТА','РЕЗУЛЬТАТ / КРИТЕРИЙ'],[
        ['8–9 недели','Проверить FAILED/PASSED, реальную дату осмотра и доступность метаданных','Словарь полей; подтверждённый момент прогноза'],
        ['10–11 недели','Получить полный журнал инспекций и локальную выборку СМР','Включены проверки без нарушений; понятна разметка'],
        ['12–13 недели','Rolling backtest; тест новых объектов; абляция служб; калибровка','Устойчивые метрики по времени и подгруппам'],
        ['14–15 недели','Согласовать цены ошибок и лимит очереди; сравнить простое дерево','Рабочий порог; precision@K; сравнение с baseline']],[2.0,5.2,4.7],row_h=.80,fonts=[('Consolas',12,True),('Arial',13,False),('Arial',12,False)])
    t(s,.7,6.03,11.9,.4,'Разработка, реализация ERP и внедрение — следующие работы проекта, вне текущей аналитической сдачи.',11,'Consolas',GRAY)
    s=tpl.page('10 — РИСКИ','Ограничения результата и способы их уменьшить')
    table(s,['РИСК','УРОВЕНЬ','КАК УМЕНЬШИТЬ'],[
        ['Архив нарушений не охватывает все инспекции','ВЫСОКИЙ','Получить полный журнал и знаменатель проверок'],
        ['Чикаго отличается от локальных процессов СМР','ВЫСОКИЙ','Независимая проверка на локальных данных'],
        ['Исторические статусы и поля могут меняться','ВЫСОКИЙ','Версии данных на момент принятия решения'],
        ['Разные даты внутри inspection_number','СРЕДНИЙ','Реальная дата осмотра; анализ чувствительности отбора'],
        ['Повторы объектов; дрейф служб и категорий','СРЕДНИЙ','Тест новых объектов; оценки по годам и службам'],
        ['Цена ошибок 5:1 пока условная','СРЕДНИЙ','Согласование; сценарии 1:1 / 3:1 / 5:1 / 10:1']],[5.3,1.7,4.9],row_h=.54,fonts=[('Arial',13,False),('Consolas',10,True),('Arial',12,False)],highlight=[(i,1) for i in range(3)])
    t(s,.7,6.08,11.9,.37,f"{ru(m['test_objects_seen_in_train_share']*100)} % объектов test уже встречались в train: текущая оценка относится к будущим проверкам известного парка.",11,'Consolas',GRAY)
    # 22–27: retain the original preparation/memory ledger, every benchmark,
    # and the complete tool-selection argument, with their original provenance.
    for idx,sec in [(4,'11 — ОБРАБОТКА · ЗАМЕРЫ 26.09.2026'),(5,'11 — ХРАНЕНИЕ · ЗАМЕРЫ 26.09.2026'),(6,'11 — PANDAS · ЗАМЕРЫ 26.09.2026'),(7,'11 — DUCKDB · ЗАМЕРЫ 26.09.2026'),(8,'11 — ВЕДОМОСТЬ · ЗАМЕРЫ 26.09.2026'),(9,'11 — ВЫБОР ИНСТРУМЕНТОВ · ЗАМЕРЫ 26.09.2026')]:
        src=tpl.ref.slides[idx];title=src.shapes[4].text
        s=tpl.page(sec,title,source=idx,historical=True)
        if idx==7:
            for sh in s.shapes:
                if sh.has_text_frame and 'ничего не загружает в pandas' in sh.text:set_text(sh,'Прямое поколоночное сканирование: DuckDB читает нужный столбец и выполняет запрос в собственной памяти; загрузка в pandas не требуется.')
        t(s,.7,6.52,7.5,.18,'Исторические измерения автора: прежний снимок и машина; не повторные замеры.',8,'Consolas',ORANGE)
    # 28: final conclusion uses the original text + four verdict cards layout.
    s=tpl.page('12 — РЕЗУЛЬТАТ АНАЛИТИЧЕСКОЙ РАБОТЫ','От данных к проверяемой точке отсчёта')
    t(s,.7,2.0,6.1,4.4,f"На полном снимке из {n(q['rows'])} записей проверены типы, полнота, уникальность, диапазоны, согласованность, метка, репрезентативность и утечка.\n\nДля модели сформировано {n(q['labeled_inspections'])} инспекций с однозначным FAILED/PASSED. Семь графиков объясняют структуру данных и ограничения анализа.\n\nНа временном тесте AP логистической модели — {ru(lr['average_precision'],3)} против {ru(dummy['average_precision'],3)} у Dummy. Но приоритет получает 97,2 % проверок: практическая полезность требует локальных данных и подтверждённого лимита внимания.",15,line=1.13)
    card(s,7.2,2.0,5.4,'Отчёт','14 СТРАНИЦ','Постановка, качество, EDA, baseline, план и риски')
    card(s,7.2,3.08,5.4,'Notebook','ВЫПОЛНЕН','Все ячейки исполнены; результаты воспроизводимы')
    card(s,7.2,4.16,5.4,'Данные','СНИМОК В GIT','Компактные таблицы + дата, источник и SHA-256')
    card(s,7.2,5.24,5.4,'Проект','САМОСТОЯТЕЛЬНЫЙ','Аналитика курсового проекта; без привязки к диплому')
    total=tpl.save();assert total==28
    print(f'Generated {total} slides using the exact reference templates.')


if __name__=='__main__':
    from src.artifacts import report_pages
    q=json.loads((ROOT/'reports/quality.json').read_text());m=json.loads((ROOT/'reports/baseline.json').read_text());eda=json.loads((ROOT/'reports/eda_insights.json').read_text());manifest=json.loads((ROOT/'data/raw/manifest.json').read_text())
    _,checks=report_pages(q,m,eda,manifest)
    make_presentation(q,m,eda,checks)
