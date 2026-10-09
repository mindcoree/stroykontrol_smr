"""Generate the report and presentation from measured pipeline artifacts."""
from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
REP = ROOT/'reports'
FIG = REP/'figures'
ORANGE, DARK, PAPER, GRAY, BORDER = 'E8591A', '1E2227', 'F2EFE8', '5B636E', 'C9C3B6'


def load(name):
    return json.loads((REP/name).read_text())


def n(x):
    return f'{int(x):,}'.replace(',', ' ')


def f(x):
    return f'{x:.3f}'


def report_pages(q,m,eda,manifest):
    test=[r for r in m['metrics'] if r['split']=='test']
    lr=next(r for r in test if r['model']=='Logistic regression')
    dummy=next(r for r in test if r['model']=='Dummy prior')
    verdict='достигнут' if m['success_criterion_met'] else 'не достигнут'
    checks=[
        ['1. Структура и типы',f"{q['columns']} полей; ошибки разбора дат: {sum(q['date_parse_failures'].values())}.",'Явные типы, строгий CSV-парсер; потеря строк не допускается.'],
        ['2. Полнота',f"Без ключей: {n(q['missing_keys'])}; без координат: {n(q['missing_coordinates'])}.",'Пропуски считаются по всем полям; координаты импутируются внутри train.'],
        ['3. Уникальность',f"Повторы ID: {n(q['duplicate_ids'])}; полные повторы: {n(q['duplicate_rows'])}.",'ID нарушения и номер инспекции имеют разную гранулярность.'],
        ['4. Допустимые значения',f"Некорректных дат: {n(q['bad_dates'])}; координат вне рамки: {n(q['outside_coordinates'])}.",'Даты 2006–09.10.2026; грубая рамка 41.6–42.1 / −88–−87.5.'],
        ['5. Согласованность',f"Инспекций с разными датами: {n(q['ambiguous_inspections'])}; обратных дат статуса: {n(q['reversed_status_dates'])}.",'Проверка единства объекта, даты, службы, категории и статуса внутри инспекции.'],
        ['6. Целевая метка',f"Бинарная выборка: {n(q['labeled_inspections'])}; FAILED: {q['failed_share']:.1%}.",'FAILED=1, PASSED=0; остальные статусы и конфликты исключены.'],
        ['7. Репрезентативность',f"Один город; {n(q['objects'])} групп объектов; 2026 неполный.",'Архив нарушений не является полным реестром всех осмотров или СМР.'],
        ['8. Утечка и время',f"Train ≤2023; validation 2024; test 2025; {len(m['features'])} исходных признаков.",'Постфактум поля исключены; версии полей на момент осмотра недоступны.'],
    ]
    pages=[]
    def page(title,*blocks):pages.append({'title':title,'blocks':list(blocks)})
    def p(s):return ('p',s)
    def h(s):return ('h',s)
    def tab(headers,rows):return ('table',(headers,rows))
    def graph(i,width=16):return ('figure',(eda[i]['id']+'.png',eda[i]['conclusion'],width))
    page('Курсовой проект · первый рубежный контроль',
         h('Анализ данных строительных инспекций для приоритизации контроля'),
         p('Тема проекта: «Автоматизация строительного контроля и управления качеством СМР». Автор: Медиханов Тимур. Аналитическая часть самостоятельного учебного проекта. Связь с дипломом не заявляется.'),
         p('Дата подготовки: 9 октября 2026 года. Источник: Chicago Building Violations, Chicago Data Portal, набор 22u3-xenr.'),
         h('Аннотация'),
         p(f"Исследован полный доступный снимок из {n(q['rows'])} записей нарушения и {q['columns']} полей. После объединения записей по номеру инспекции и отбора однозначных результатов получено {n(q['labeled_inspections'])} наблюдений для бинарной классификации. Проведены аудит качества по восьми пунктам, разведочный анализ с семью графиками и сравнение константной модели с логистической регрессией на временном тесте."),
         p(f"На 2025 годе Average Precision логистической модели составляет {f(lr['average_precision'])}; у константной точки отсчёта — {f(dummy['average_precision'])}. Предварительный критерий успеха {verdict}. Результат относится к зарегистрированным в архиве инспекциям и требует проверки на локальных данных перед практическим использованием."),
         h('Состав сдаваемого результата'),
         p('Постановка задачи; паспорт источника; восемь проверок качества; EDA с выводами; baseline; план дальнейшей работы и риски. Код, notebook, таблицы, графики и материалы размещены в репозитории. Разработка ERP-приложения и внедрение не входят в оценку этой работы.'))
    page('1. Постановка задачи и критерий успеха',
         p('Пользователь результата — руководитель службы строительного контроля или координатор инспекций. Его решение — какие из уже запланированных проверок требуют более раннего внимания и дополнительной подготовки. Выход аналитики — вероятность FAILED и очередь приоритетов; решение о результатах самой проверки принимает инспектор.'),
         p('Единица наблюдения — инспекция, определённая inspection_number. Цель — оценить P(FAILED | метаданные запланированной инспекции) среди наблюдений с однозначным FAILED или PASSED, представленных в журнале нарушений. Задача является бинарной классификацией с последующим ранжированием.'),
         p('Предполагаемый момент прогноза — до осмотра. На входе: назначенная служба, категория проверки, календарная дата и известные координаты объекта. Их доступность в реальной системе пока не подтверждена. Текущий статус, список нарушений и результат устранения не используются.'),
         tab(['Ошибка','Практическое последствие','Учебная цена'],[
             ['FN: FAILED предсказан как PASSED','Проблемная проверка получает низкий приоритет; задержка реакции и дополнительной подготовки.','5 условных единиц'],
             ['FP: PASSED предсказан как FAILED','Лишнее внимание и расход ограниченного времени специалиста.','1 условная единица']]),
         p('Соотношение 5:1 — учебная гипотеза, а не измеренный денежный ущерб и не экспертная оценка. Оно применяется одинаково к обеим моделям. Без подтверждения стоимости и доступной мощности проверок оптимальный порог нельзя считать операционным решением.'),
         p('Предварительный критерий: на отложенном годе AP выше константного baseline минимум на 0.05, recall FAILED не ниже 0.80, precision не ниже доли FAILED на этом тесте. Порог выбирается только на validation по минимуму (5×FN+FP)/N. AP оценивает ранжирование, recall — пропуск проблемных проверок, precision — расход внимания.'))
    page('2. Источник, формат и ограничения доступа',
         p('Первичный источник — Department of Buildings города Чикаго. Владелец публикует исторические записи нарушений с 2006 года; несколько нарушений могут относиться к одной инспекции. Данные носят информационный характер и могут не отражать текущее состояние здания [1].'),
         tab(['Параметр','Значение'],[
             ['Источник','data.cityofchicago.org · 22u3-xenr'],['Строк / полей',f"{n(q['rows'])} / {q['columns']}"],
             ['Период событий',f"{str(q['date_min'])[:10]} — {str(q['date_max'])[:10]}"],
             ['Групп объектов / номеров инспекции',f"{n(q['objects'])} / {n(q['inspection_ids'])}"],
             ['Различных кодов нарушения',n(q['violation_types'])],
             ['CSV / Parquet ZSTD',f"{q['csv_mb']:.1f} / {q['parquet_mb']:.1f} МБ"],
             ['Скачано (Asia/Almaty)',manifest['retrieved_at'][:19].replace('T',' ')]]),
         p('Формат исходного экспорта — CSV UTF-8, 26 полей, даты представлены текстом. Для вычислений используется Parquet с явными типами и сжатием ZSTD. Открытый CSV экспорт скачан без учётной записи и без API-токена. Интернет требуется для получения нового снимка; после загрузки расчёты выполняются локально.'),
         p('Ключевые поля: ID — запись нарушения; inspection_number — связь записей проверки; property_group — группа объектов; violation_date — дата события, используемая как приближение даты инспекции; inspection_status — наблюдаемый результат; department_bureau и inspection_category — организационный контекст; latitude/longitude — география.'),
         p('Ограничения: источник изменяется; сервер может ограничить частоту запросов, а исторические статусы могут обновляться. Для повторения чисел сохранены SHA-256, время загрузки, метаданные и компактная таблица инспекций. Адреса и комментарии не нужны модели и не включены в компактную таблицу. Открытость источника не означает наличие подтверждённой свободной лицензии: новая лицензия на чужие данные не назначается.'))
    page('3. Качество данных: пункты 1–4',
         p('Чек-лист преподавателя в предоставленном задании не раскрыт. Ниже применён стандартный аналитический чек-лист из восьми пунктов; это явное допущение проекта. Все количественные проверки выполнены на полном загруженном CSV, без подвыборки. Подробная машинная сводка — reports/quality.json.'),
         tab(['Проверка','Измеренный результат','Решение'],checks[:4]),
         h('Полнота и преобразование типов'),
         p('Пропуском считается NULL или пустая строка после удаления крайних пробелов. Даты разбираются по явным форматам; неудачные преобразования отдельно проверяются против исходного текста. CSV читается в строгом режиме: ignore_errors не используется, поэтому ошибочные строки не могут исчезнуть незаметно.'),
         p('Для признаков модели числовые пропуски заполняются медианой только обучающего периода, одновременно создаётся индикатор отсутствия значения. Категориальные пропуски обозначаются UNKNOWN. Пустые даты и ключи не имитируются искусственными значениями.'),
         h('Повторы и географические границы'),
         p('Совпадение inspection_number ожидаемо: это связь нескольких нарушений одной проверки, а не дубликат строки. При повторе ID сохраняется запись с самой поздней датой изменения. Координаты вне грубой рамки Чикаго заменяются пропуском, а сама инспекция сохраняется. Рамка — проверка здравого смысла, а не официальный полигон административной границы.'))
    page('3. Качество данных: пункты 5–8',
         tab(['Проверка','Измеренный результат','Решение'],checks[4:]),
         p('Конфликт определяется как более одного непустого значения статуса, объекта, даты, службы или категории внутри inspection_number. Такие инспекции не используются для обучения: выбор первого или последнего значения без правила бизнес-процесса создавал бы произвольную разметку. Для EDA они сохранены с признаком ambiguous.'),
         p(f"Все {n(q['ambiguous_inspections'])} исключения связаны с разными violation_date внутри одной инспекции; по статусу, объекту, службе и категории противоречий не обнаружено. Это неоднозначность приближения даты, а не доказанная ошибка источника. Консервативный отбор может сместить выборку: чувствительность к этому правилу нужно проверить после получения реальной даты осмотра."),
         p('FAILED и PASSED трактуются как наблюдаемые классы. CLOSED, HOLD и UNKNOWN не считаются успешной проверкой. Их число сохраняется в сводке, чтобы показать цену фильтрации. Условия отбора определены до обучения и одинаковы для всех периодов.'),
         p('Репрезентативность ограничена механизмом попадания в журнал. Полные данные о зданиях без зарегистрированных нарушений и полный журнал всех инспекций отсутствуют. Поэтому нельзя оценить общий риск любого здания города или автоматически переносить частоты на стройплощадки Казахстана.'),
         p('Утечка ограничена исключением текущих нарушений, комментариев, дат устранения, inspection_waived и результата проверки из признаков. Исторические target-агрегаты из прежнего notebook также исключены: последовательный порядок инспекций одного дня не подтверждает доступность их результатов. Остаточный риск — один современный снимок без версий полей на дату прогноза.'))
    page('4. Подготовка данных и аналитический протокол',
         tab(['Этап','Что выполняется','Выход'],[
             ['raw','Загрузка полного CSV; хеш; метаданные; строгий разбор; типизированный Parquet.','CSV и manifest.json'],
             ['processed','Нормализация строк; дедупликация ID; проверка дат; агрегация; поиск конфликтов.','inspections.parquet'],
             ['model data','Отбор FAILED/PASSED без конфликтов; календарные признаки; сохранение полей для аудита.','model_data.parquet'],
             ['analysis','7 EDA-графиков; временное разделение; fit train; порог validation; оценка test.','JSON, CSV, PNG, SVG']]),
         p(f"Из {n(q['rows'])} строк сформировано {n(q['aggregated_inspections'])} инспекций с корректным номером и датой. Для бинарной задачи осталось {n(q['labeled_inspections'])}; исключено {n(q['aggregated_inspections']-q['labeled_inspections'])} из-за неоднозначного статуса или конфликтующих атрибутов. По каждой инспекции сохранено число исходных записей."),
         p('Дата инспекции приближена минимальной violation_date. Проверка единства даты внутри группы помогает обнаружить сомнительное соответствие, но не доказывает, что дата нарушения равна фактическому времени осмотра. Это ограничение должно быть снято полным реестром инспекций.'),
         p('Поля n_violations и n_codes применяются в разведочном анализе, поскольку отражают структуру регистрации. Они отсутствуют в списке входов baseline: до текущего осмотра неизвестно, сколько нарушений будет найдено. Адрес, ID нарушения и инспектор тоже исключены из модели.'),
         p('Код разделяет статистику по нарушениям и по инспекциям. Общие EDA-графики используют весь снимок; выбор порога и обучение никогда не используют test. EDA всего архива считается описанием данных, а не независимой проверкой гипотез.'))
    page('5. EDA: полнота и временной охват',graph(0),graph(1),
         p('Аналитический вывод: большой размер архива не устраняет систематические пропуски и изменение процесса регистрации. Доля отсутствующих значений измеряется отдельно от аномальных диапазонов; неполный последний год не сравнивается с полными годами как равный по экспозиции.'))
    page('5. EDA: метка и единица наблюдения',graph(2),graph(3),
         p('Аналитический вывод: классифицировать следует инспекции, а не записи нарушения. Иначе сложные проверки получили бы больший вес, а записи одной проверки могли бы попасть одновременно в обучение и тест. Остальные статусы требуют отдельного исследования, а не автоматического превращения в отрицательный класс.'))
    page('5. EDA: различия служб',graph(4),
         h('Интерпретация организационного контекста'),
         p('Сравниваются десять служб с наибольшим количеством наблюдений в бинарной выборке. На графике показаны доли FAILED и 95% интервалы Уилсона. Большие группы дают более узкий описательный интервал, но повторные проверки одного объекта нарушают независимость наблюдений.'),
         p('Различия служб могут объясняться содержанием процедур, составом объектов и тем, какие результаты попадают в журнал. Почти постоянный класс у отдельной службы — повод проверить семантику учёта; это не доказательство высокого или низкого качества работы службы.'),
         p('Служба и категория включены в простую модель как доступные метаданные. Следующая проверка — модель без этих полей, чтобы измерить, какая часть качества является следствием маршрутизации инспекций. Перенос между службами и городами должен оцениваться отдельно.'),
         p('Таблица bureau_rates.csv содержит число инспекций и долю FAILED для каждого отображённого подразделения. Она позволяет проверить, что вывод не основан на редкой группе из нескольких событий.'))
    page('5. EDA: сезонность и география',graph(5),graph(6),
         p('Аналитический вывод: календарные и пространственные поля могут описывать состав и интенсивность проверок. Для оценки районного риска нужны внешние знаменатели: количество зданий, их возраст и полный реестр осмотров. В проекте причинные и территориальные выводы не заявляются.'))
    page('6. Baseline: признаки, модели и разделение',
         tab(['Выборка','Период','Инспекции','FAILED'],[[k,v['date_from']+' — '+v['date_to'],n(v['n']),f"{v['failed_share']:.1%}"] for k,v in m['splits'].items()]),
         p(f"Наблюдения 2026 года ({n(m['excluded_2026'])}) не участвуют в оценке модели: год неполный. Разделение календарное, случайное перемешивание не применяется. Повторные объекты между train и test разрешены, поскольку предполагается работа на последующих проверках известного парка объектов."),
         p('Точка отсчёта — DummyClassifier(strategy="prior"). Она оценивает частоту FAILED по train и выдаёт одинаковую вероятность всем инспекциям; признаки не использует. Такая модель выявляет, насколько результат объясняется одной базовой частотой [2].'),
         p('Вторая простая модель — логистическая регрессия с L2-регуляризацией (C=1, lbfgs, max_iter=2000). Два категориальных входа кодируются OneHotEncoder(handle_unknown="ignore"); шесть числовых входов заполняются медианой и стандартизуются. Преобразования обучаются в Pipeline только на train [3].'),
         p('Входы: department_bureau, inspection_category, year, month_sin, month_cos, dayofweek, latitude, longitude. Циклическая пара sin/cos отражает близость декабря и января. Год описывает линейный тренд, поэтому экстраполяция на будущие годы — отдельный риск.'),
         p('Для каждой модели перебирается порог от 0 до 1 с шагом 0.005 на 2024 годе. Выбирается минимум учебной стоимости 5×FN+FP; при равенстве берётся меньший порог. После этого модель и порог фиксируются и один раз оцениваются на 2025 годе. Модель не переобучается на validation.'))
    page('6. Baseline: результаты на 2025 годе',
         tab(['Модель','AP','ROC-AUC','Precision','Recall','Cost/N'],[[r['model'],f(r['average_precision']),f(r['roc_auc']),f(r['precision']),f(r['recall']),f(r['cost_per_inspection'])] for r in test]),
         ('figure',('baseline_confusion.png','Матрицы ошибок: строки — факт, столбцы — прогноз; положительный класс FAILED.',16)),
         p(f"Порог логистической модели — {lr['threshold']:.3f}. На тесте: FN={n(lr['fn'])}, FP={n(lr['fp'])}, TP={n(lr['tp'])}, TN={n(lr['tn'])}. Balanced accuracy={lr['balanced_accuracy']:.3f}, F1={lr['f1']:.3f}, Brier={lr['brier']:.3f}. Предварительный критерий успеха {verdict}; это оценка внутри доступного архива."),
         p(f"При данном пороге приоритет получают {(lr['tp']+lr['fp'])/lr['n']:.1%} тестовых инспекций; доля ложных тревог среди PASSED — {lr['fp']/(lr['fp']+lr['tn']):.1%}. Учебная стоимость ниже Dummy на {(1-lr['cost_per_inspection']/dummy['cost_per_inspection']):.1%}. Высокий recall достигнут широкой очередью, поэтому практическая экономия внимания пока не доказана."),
         p('AP здесь — average_precision_score, сумма precision, взвешенная приростами recall, а не трапецеидальная площадь PR-кривой [4]. ROC-AUC дополнительно характеризует ранжирование. Brier оценивает средний квадрат ошибки вероятности; меньше — лучше.'),
         p(f"Доля объектов теста, встречавшихся в train: {m['test_objects_seen_in_train_share']:.1%}. Такой тест соответствует следующим проверкам существующего парка, но не подтверждает перенос на новые объекты. Для него нужна дополнительная оценка по группам property_group."))
    page('7. План дальнейшей работы и риски',
         tab(['Этап после РК-1','Результат','Критерий завершения'],[
             ['Недели 8–9','Проверить семантику FAILED/PASSED, даты и доступность метаданных.','Словарь полей и подтверждённый момент прогноза.'],
             ['Недели 10–11','Получить полный реестр инспекций и локальную выборку СМР.','Есть проверки без нарушений и однозначная разметка.'],
             ['Недели 12–13','Rolling backtest; тест новых объектов; абляция служб; калибровка.','Стабильные метрики по периодам и подгруппам.'],
             ['Недели 14–15','Проверить стоимость ошибок и лимит очереди; сравнить простое дерево.','Порог и precision@K согласованы с пользователем.']]),
         tab(['Риск','Значимость','Способ уменьшить'],[
             ['Смещение журнала нарушений','Высокая','Получить полный знаменатель инспекций; ограничить выводы до этого.'],
             ['Чикаго ≠ локальный контроль СМР','Высокая','Независимая локальная оценка; не переносить готовый порог.'],
             ['Обновление исторических статусов','Высокая','Версии событий и данные на момент принятия решения.'],
             ['Дрейф служб и категорий','Средняя','Оценка по годам и службам; мониторинг частот.'],
             ['Неизвестная цена ошибки','Средняя','Экспертное согласование, анализ 1:1 / 3:1 / 5:1 / 10:1.'],
             ['Повторы объектов и редкие группы','Средняя','Групповые оценки; интервалы с учётом кластеров.']]),
         p('В рамках предмета дальнейшие шаги остаются аналитическими: уточнение данных, качества, постановки и простой модели. Разработка приложения, интеграция с ERP и внедрение — отдельные будущие работы проекта, не часть текущей сдачи.'))
    page('8. Воспроизводимость, выводы и источники',
         p('Полный запуск: создать Python-окружение, установить requirements.txt, выполнить python -m src.pipeline и python -m src.artifacts. Интерактивный вариант — notebooks/course_project.ipynb. Для повторения baseline на зафиксированном снимке интернет не требуется: компактные inspections.parquet и model_data.parquet включены в репозиторий.'),
         p('Исходный CSV и raw Parquet велики и исключены из Git; код скачивает их из первоисточника. Повторная загрузка живого архива может дать другие числа. SHA-256 фиксирует идентичность исходного файла, но не гарантирует его вечную доступность на сервере. Версии использованных библиотек сохранены в requirements.txt и reports/environment.json.'),
         p('SHA-256 исходного CSV: '+manifest['sha256']),
         p('Старые замеры производительности от 26.09.2026 сохранены отдельно как историческое приложение: CSV 930.39 МБ, Parquet Snappy 149.29 МБ; чтение 15.25 и 1.581 секунды. Они относятся к иной машине и прежнему снимку, не выдаются за повторные замеры в этом отчёте. При выбранном масштабе Parquet и DuckDB подходят для локального аудита и агрегации, pandas — для таблицы инспекций.'),
         h('Выводы'),
         p(f"Аналитическая задача сформулирована условно на охват источника; качество проверено, а гранулярность исправлена до уровня инспекции. Семь EDA-графиков показывают пропуски, временной охват, баланс метки, множественность нарушений, службы, сезон и географию. Простая модель улучшает AP на {lr['average_precision']-dummy['average_precision']:.3f} относительно частотной точки отсчёта. Практическая переносимость остаётся открытой до получения локальных данных и истории версий."),
         h('Источники (проверены 09.10.2026)'),
         p('[1] Chicago Data Portal. Building Violations. https://data.cityofchicago.org/Buildings/Building-Violations/22u3-xenr ; метаданные: data/raw/source_metadata.json.'),
         p('[2] scikit-learn. DummyClassifier. https://scikit-learn.org/stable/modules/generated/sklearn.dummy.DummyClassifier.html'),
         p('[3] scikit-learn. Common pitfalls: leakage and preprocessing. https://scikit-learn.org/stable/common_pitfalls.html'),
         p('[4] scikit-learn. average_precision_score. https://scikit-learn.org/stable/modules/generated/sklearn.metrics.average_precision_score.html'))
    assert len(pages)==14
    return pages,checks


def make_report(pages):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(2)
    style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=DPt(11)
    style.paragraph_format.space_after=DPt(6);style.paragraph_format.line_spacing=1.08
    for name,size in [('Heading 1',17),('Heading 2',13)]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=DPt(size);s.font.color.rgb=DColor.from_string(DARK)
    header=sec.header.paragraphs[0];header.text='СТРОЙКОНТРОЛЬ СМР  /  АНАЛИТИЧЕСКИЙ ПРОЕКТ  /  РК-1'
    header.style='Caption';header.runs[0].font.size=DPt(8)
    foot=sec.footer.paragraphs[0];foot.text='Медиханов Тимур · 09.10.2026                                      '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    markdown=[]
    for i,page in enumerate(pages):
        if i:doc.add_page_break()
        doc.add_heading(page['title'],level=1);markdown.append('# '+page['title']+'\n')
        for kind,val in page['blocks']:
            if kind in ['p','h']:
                doc.add_heading(val,level=2) if kind=='h' else doc.add_paragraph(val)
                markdown.append(('## ' if kind=='h' else '')+val+'\n')
            elif kind=='table':
                headers,rows=val;t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
                for cell,text in zip(t.rows[0].cells,headers):cell.text=text
                for row in rows:
                    for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
                for row in t.rows:
                    trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                    for cell in row.cells:
                        for par in cell.paragraphs:
                            par.paragraph_format.space_after=DPt(3)
                            for r in par.runs:r.font.size=DPt(9)
                doc.add_paragraph().paragraph_format.space_after=DPt(0)
                markdown+=['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
                markdown+=['| '+' | '.join(str(x).replace('|','/') for x in row)+' |' for row in rows];markdown.append('')
            elif kind=='figure':
                filename,caption,width=val
                # Two-figure pages use compact charts; standalone pages use wider charts.
                two=sum(k=='figure' for k,v in page['blocks'])==2
                doc.add_picture(str(FIG/filename),width=Cm(15.3 if two else width))
                par=doc.add_paragraph(caption);par.paragraph_format.space_after=DPt(8)
                for r in par.runs:r.font.size=DPt(10)
                markdown.append(f'![{caption}](figures/{filename})\n\n{caption}\n')
    doc.save(REP/'course_report.docx')
    (REP/'course_report.md').write_text('\n'.join(markdown),encoding='utf-8')


def text(slide,x,y,w,h,s,size=16,color=DARK,bold=False,font='Arial'):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(str(s).split('\n')):
        par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        par.text=line;par.font.name=font;par.font.size=Pt(size);par.font.bold=bold;par.font.color.rgb=RGBColor.from_string(color)
        par.space_after=Pt(7)
    return sh


def rect(slide,x,y,w,h,fill,line=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(.6)
    else:sh.line.fill.background()
    return sh


def base(deck,section,title):
    s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(PAPER)
    border=rect(s,.3,.3,12.73,6.9,PAPER,BORDER)
    rect(s,.7,.67,.13,.13,ORANGE)
    text(s,.95,.6,11,.35,section,11,ORANGE,True)
    text(s,.7,.95,11.9,.8,title,28,DARK,True,'Bahnschrift')
    rect(s,.7,1.72,11.85,.012,BORDER)
    text(s,.7,6.84,7.4,.3,'Chicago Building Violations · снимок 09.10.2026 · учебный проект',8,GRAY)
    rect(s,8.45,6.70,4.55,.5,PAPER,DARK)
    text(s,8.60,6.76,2.8,.17,'СТРОЙКОНТРОЛЬ СМР · РК-1',7,GRAY,True)
    text(s,8.60,6.96,2.8,.20,'Медиханов Тимур',10,DARK,True)
    text(s,11.55,6.76,1.2,.35,str(len(deck.slides)),14,ORANGE,True)
    return s


def slide_table(s,headers,rows,widths=None,y=2.1,font=15):
    widths=widths or [11.85/len(headers)]*len(headers)
    x=.7
    for header,w in zip(headers,widths):text(s,x,y,w-.12,.35,header,11,GRAY,True);x+=w
    y+=.55
    for row in rows:
        rect(s,.7,y-.08,11.85,.012,BORDER);x=.7
        height=.70 if len(rows)<=4 else .49
        for v,w in zip(row,widths):text(s,x,y,w-.14,height,str(v),font,DARK);x+=w
        y+=height+.12


def copy_slide(deck,src):
    s=deck.slides.add_slide(deck.slide_layouts[6])
    for sh in src.shapes:s.shapes._spTree.insert_element_before(deepcopy(sh.element),'p:extLst')
    return s


def make_presentation(q,m,eda,checks):
    original=Presentation(ROOT/'archive/presentation_original.pptx')
    deck=Presentation();deck.slide_width=original.slide_width;deck.slide_height=original.slide_height
    cover=copy_slide(deck,original.slides[0])
    replacements={
        '0,9 ГБ':f"{q['rows']/1e6:.2f} млн",
        'ЗАДАЧА 4 · РК-1 · DATA SCIENCE':'КУРСОВОЙ ПРОЕКТ · РК-1 · DATA SCIENCE',
        'Приложение для корпоративной ERP-платформы':'Аналитическая часть самостоятельного учебного проекта',
        'Обоснование программных средств на полном архиве\npandas  ·  Parquet  ·  DuckDB':'Постановка задачи · качество данных · 7 графиков EDA\nBaseline · план дальнейшей работы · воспроизводимый код',
        '2\u00a0030\u00a0101 запись\n930 МБ в CSV':f"{n(q['rows'])} записей\n{q['csv_mb']:.0f} МБ в CSV",
        'ДАТА ЗАМЕРОВ':'ДАТА АНАЛИЗА','26.09.2026 12:17':'09.10.2026'}
    for sh in cover.shapes:
        if sh.has_text_frame and sh.text in replacements:
            value=replacements[sh.text];style=sh.text_frame.paragraphs[0].runs[0] if sh.text_frame.paragraphs[0].runs else None
            props=deepcopy(style._r.rPr) if style is not None and style._r.rPr is not None else None
            sh.text_frame.clear()
            for i,line in enumerate(value.split('\n')):
                p=sh.text_frame.paragraphs[0] if i==0 else sh.text_frame.add_paragraph();r=p.add_run();r.text=line
                if props is not None:r._r.insert(0,deepcopy(props))
    s=base(deck,'01 — ПОСТАНОВКА ЗАДАЧИ','Кому и для какого решения нужен результат')
    for x,label,body in [(.7,'ПОЛЬЗОВАТЕЛЬ','Координатор инспекций и руководитель строительного контроля.'),(4.8,'РЕШЕНИЕ','Приоритет внимания к уже запланированной инспекции.'),(8.8,'ВЫХОД','Вероятность FAILED среди записанных FAILED/PASSED.')]:
        text(s,x,2.1,3.7,.4,label,11,ORANGE,True);text(s,x,2.65,3.7,1.4,body,18)
    text(s,.7,4.2,11.7,.75,'FN: пропущенная проблемная проверка — 5 у.е.\nFP: лишнее внимание к успешной проверке — 1 у.е.',18)
    text(s,.7,5.5,11.7,.85,'Критерий: AP ≥ Dummy + 0.05; recall ≥ 0.80; precision ≥ доли FAILED.\nЦена ошибки — учебная гипотеза. Решение проверяет специалист.',16,GRAY)
    s=base(deck,'02 — ИСТОЧНИК ДАННЫХ','Полный архив, но ограниченный охват')
    slide_table(s,['ПАРАМЕТР','СНИМОК 09.10.2026'],[
        ['Источник','Chicago Data Portal · 22u3-xenr'],['Строк / столбцов',f"{n(q['rows'])} / {q['columns']}"],
        ['Инспекции после агрегации',n(q['aggregated_inspections'])],['Бинарная выборка',n(q['labeled_inspections'])],
        ['Формат / доступ','CSV → Parquet ZSTD · открытый экспорт без токена']],widths=[4.1,7.75])
    text(s,.7,5.9,11.8,.6,'Журнал нарушений Чикаго: не все инспекции и не локальные данные СМР.\nРезультат нельзя переносить на другой город без проверки.',15,ORANGE,True)
    for offset in [0,4]:
        s=base(deck,f'03 — КАЧЕСТВО ДАННЫХ · {offset+1}–{offset+4}','Восемь проверок качества данных')
        slide_table(s,['ПРОВЕРКА','РЕЗУЛЬТАТ','РЕШЕНИЕ'],checks[offset:offset+4],widths=[2.5,4.2,5.15],font=13)
        text(s,.7,6.0,11.8,.5,'Использован стандартный чек-лист: перечень преподавателя не предоставлен.',12,GRAY)
    s=base(deck,'04 — ПОДГОТОВКА','Одна строка модели — одна инспекция')
    for x,label,value,body in [(.7,'RAW',n(q['rows']),'Записи нарушений\nСтрогий CSV-парсер\nХеш и метаданные'),(4.8,'PROCESSED',n(q['aggregated_inspections']),'Группировка inspection_number\nПроверка конфликтов\nЧисло нарушений для EDA'),(8.8,'MODEL DATA',n(q['labeled_inspections']),'Только FAILED / PASSED\nКонфликты исключены\n8 входных признаков')]:
        rect(s,x,2.1,3.75,3.6,'FFFFFF',BORDER);text(s,x+.2,2.35,3.35,.4,label,12,ORANGE,True)
        text(s,x+.2,3.0,3.35,.65,value,27,DARK,True,'Bahnschrift');text(s,x+.2,4.0,3.35,1.4,body,16)
    text(s,.7,6.0,11.8,.5,'Текущие нарушения, статус и результаты устранения исключены из входов модели.',15,GRAY)
    for i,ins in enumerate(eda):
        s=base(deck,f'05 — РАЗВЕДОЧНЫЙ АНАЛИЗ · {i+1}/7',ins['title'][3:])
        image=s.shapes.add_picture(str(FIG/(ins['id']+'.png')),Inches(.55),Inches(2.05),width=Inches(8.55))
        text(s,9.35,2.1,3.0,.4,'ВЫВОД',11,ORANGE,True)
        text(s,9.35,2.65,2.95,3.65,ins['conclusion'],15)
    s=base(deck,'06 — BASELINE','Временная оценка без текущих нарушений')
    slide_table(s,['ВЫБОРКА','ПЕРИОД','ИНСПЕКЦИИ','FAILED'],[[k,v['date_from'][:4]+'–'+v['date_to'][:4],n(v['n']),f"{v['failed_share']:.1%}"] for k,v in m['splits'].items()],widths=[3,3,3,2.85])
    text(s,.7,5.0,11.8,1.3,'Dummy prior → частота FAILED по train\nLogistic regression → служба, категория, календарь и координаты\nПорог: минимум 5×FN + FP на 2024; 2025 используется только для оценки.',17)
    s=base(deck,'06 — BASELINE · РЕЗУЛЬТАТ','Простая модель против частотной точки отсчёта')
    test=[r for r in m['metrics'] if r['split']=='test']
    slide_table(s,['МОДЕЛЬ','AP','ROC-AUC','PRECISION','RECALL','COST/N'],[[r['model'],f(r['average_precision']),f(r['roc_auc']),f(r['precision']),f(r['recall']),f(r['cost_per_inspection'])] for r in test],widths=[3.7,1.4,1.65,1.8,1.55,1.75],font=15)
    s.shapes.add_picture(str(FIG/'baseline_confusion.png'),Inches(.7),Inches(4.15),width=Inches(7.0))
    lr=next(r for r in test if r['model']=='Logistic regression')
    text(s,8.05,4.15,4.15,2.2,f"Порог: {lr['threshold']:.3f}\nFN: {n(lr['fn'])} · FP: {n(lr['fp'])}\nКритерий: {'достигнут' if m['success_criterion_met'] else 'не достигнут'}\nВ приоритете {(lr['tp']+lr['fp'])/lr['n']:.1%} проверок.\nЭкономия внимания ещё не доказана.",16,ORANGE,True)
    s=base(deck,'07 — ПЛАН И РИСКИ','Что проверить после первого рубежного контроля')
    slide_table(s,['ЭТАП','РАБОТА','ОСНОВНОЙ РИСК'],[
        ['8–9 недели','Уточнить статусы, даты и доступность полей','Утечка через изменённые поля'],
        ['10–11 недели','Получить полный журнал и локальные данные','Смещение архива нарушений'],
        ['12–13 недели','Backtest, новые объекты, модель без служб','Дрейф и переносимость'],
        ['14–15 недели','Согласовать стоимость и лимит очереди','Учебная цена ≠ реальный ущерб']],widths=[2.1,5.4,4.35],font=15)
    text(s,.7,6.0,11.8,.45,'Текущая сдача — аналитика. Разработка и внедрение ERP остаются вне её объёма.',15,GRAY)
    s=base(deck,'08 — РЕЗУЛЬТАТ ПРОЕКТА','Все шесть пунктов задания покрыты')
    text(s,.7,2.1,7.5,3.9,'01  Постановка: пользователь, решение, критерий, ошибки\n02  Источник: полный снимок, формат, доступ и ограничения\n03  Качество: восемь проверок с измеренными результатами\n04  EDA: семь графиков и выводы по каждому\n05  Baseline: Dummy prior и логистическая регрессия\n06  План: этапы, риски и способы их уменьшить',20)
    rect(s,8.8,2.1,3.75,3.8,DARK)
    text(s,9.05,2.45,3.25,.5,'ВОСПРОИЗВОДИМОСТЬ',11,ORANGE,True)
    text(s,9.05,3.15,3.25,2.2,'Отчёт · notebook · код\nЗафиксированный снимок\nМетрики и графики\nGitHub-репозиторий',18,'FFFFFF')
    text(s,.7,6.1,11.8,.4,'Ограничение: качество модели подтверждено только внутри доступного архива.',14,GRAY)
    for idx in [5,6,7,8]:
        s=copy_slide(deck,original.slides[idx])
        # Preserve the author's original benchmark designs and measured values.
        for sh in s.shapes:
            if sh.has_text_frame and sh.text.startswith('0') and '—' in sh.text:
                sh.text='ПРИЛОЖЕНИЕ · ИСХОДНЫЕ ЗАМЕРЫ 26.09.2026'
                for p in sh.text_frame.paragraphs:
                    p.font.name='Arial';p.font.size=Pt(11);p.font.color.rgb=RGBColor.from_string(ORANGE)
        text(s,.7,6.5,7.45,.20,'Иная машина и прежний снимок; показатели не измерены повторно.',9,ORANGE,True)
        # Correct technical overstatement about memory mapping in the old slide.
        for sh in s.shapes:
            if sh.has_text_frame and 'ничего не загружает в pandas' in sh.text:
                sh.text='DuckDB читает нужный столбец из Parquet и использует собственную память для выполнения запроса.'
                for p in sh.text_frame.paragraphs:p.font.name='Arial';p.font.size=Pt(13)
        nums=[sh for sh in s.shapes if sh.has_text_frame and sh.text in ['6','7','8','9','10'] and sh.left>Inches(11)]
        for sh in nums:
            old=sh.text;sh.text=str(len(deck.slides)) if old!='10' else '21'
            for p in sh.text_frame.paragraphs:p.font.size=Pt(13);p.font.name='Arial'
    assert len(deck.slides)==21
    deck.save(ROOT/'presentation.pptx')


def main():
    q,m,eda=load('quality.json'),load('baseline.json'),load('eda_insights.json')
    manifest=json.loads((ROOT/'data/raw/manifest.json').read_text())
    pages,checks=report_pages(q,m,eda,manifest)
    make_report(pages);make_presentation(q,m,eda,checks)
    print('Generated reports/course_report.docx, .md and presentation.pptx')


if __name__=='__main__':main()
